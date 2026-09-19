"""Recompute totals from the source xlsx and confirm they match the Google Sheet claims.

Used by Phase 6 (fix-verification pass) of the audit-response-packet-review skill.

Usage:
    python3 verify_xlsx_totals.py \\
        --xlsx "/tmp/jorge-auditor/Jeffrey Nicholas Lee-Chan - 2025 Year Examination.xlsx" \\
        --sheet "Venmo Transactions" \\
        --year 2024 \\
        --payee-column 2 \\
        --amount-column 1 \\
        --date-column 0 \\
        --memo-column 2 \\
        --expected-total 19294.00

Notes:
- Requires openpyxl (`pip install openpyxl`).
- Excel date serials are converted via the 1900-base epoch (1899-12-30).
- The "memo classifier" is intentionally loose — the point is to show the
  grand total matches the claim, not to perfectly bucket every row.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
except ImportError:
    print("openpyxl not installed; pip install openpyxl", file=sys.stderr)
    sys.exit(2)


def excel_date_to_dt(serial: int) -> datetime:
    return datetime(1899, 12, 30) + timedelta(days=int(serial))


def classify(memo: str) -> str:
    m = memo.lower()
    if "advance" in m:
        return "advance"
    if "bonus" in m:
        return "bonus"
    if re.search(r"\d+(\.\d+)?\s*hrs", m) or re.search(r"\d+/\d+/\d+\s+\d+(\.\d+)?\s*hrs", m):
        return "wages_hours"
    if "declined" in m:
        return "declined"
    return "reimb_or_other"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--date-column", type=int, default=0)
    ap.add_argument("--amount-column", type=int, default=1)
    ap.add_argument("--memo-column", type=int, default=2)
    ap.add_argument("--payee-column", type=int, default=-1,
                    help="If >=0, also filter rows where this column contains --payee-match")
    ap.add_argument("--payee-match", default="")
    ap.add_argument("--expected-total", type=float, default=None,
                    help="If set, exit non-zero unless grand sum matches to the cent.")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True, read_only=True)
    if args.sheet not in wb.sheetnames:
        print(f"sheet {args.sheet!r} not in {wb.sheetnames}", file=sys.stderr)
        return 2
    ws = wb[args.sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("empty sheet", file=sys.stderr)
        return 2

    buckets: dict[str, float] = {}
    grand = 0.0
    n = 0
    for row in rows[1:]:
        if not row or row[args.date_column] is None or row[args.amount_column] is None:
            continue
        if args.payee_column >= 0:
            payee = str(row[args.payee_column] or "")
            if args.payee_match.lower() not in payee.lower():
                continue
        pay_date = row[args.date_column]
        amount = row[args.amount_column]
        memo = row[args.memo_column] if args.memo_column < len(row) else ""
        if isinstance(pay_date, datetime):
            dt = pay_date
        elif isinstance(pay_date, (int, float)):
            dt = excel_date_to_dt(int(pay_date))
        else:
            continue
        if dt.year != args.year:
            continue
        try:
            amt = float(amount)
        except (TypeError, ValueError):
            continue
        bucket = classify(str(memo))
        buckets[bucket] = buckets.get(bucket, 0.0) + amt
        grand += amt
        n += 1

    print(f"year={args.year}  rows={n}  grand_total=${grand:,.2f}")
    for bucket, total in sorted(buckets.items(), key=lambda kv: -kv[1]):
        print(f"  {bucket:<18} ${total:>12,.2f}")
    wages = buckets.get("wages_hours", 0.0) + buckets.get("bonus", 0.0)
    reimb = buckets.get("reimb_or_other", 0.0) + buckets.get("declined", 0.0)
    print(f"  wages (hrs+bonus)  ${wages:>12,.2f}")
    print(f"  reimb+declined     ${reimb:>12,.2f}")
    print(f"  advances           ${buckets.get('advance', 0.0):>12,.2f}")

    if args.expected_total is not None:
        diff = abs(grand - args.expected_total)
        if diff > 0.005:
            print(f"MISMATCH: expected ${args.expected_total:,.2f} got ${grand:,.2f} (diff ${diff:.2f})", file=sys.stderr)
            return 1
        print(f"OK: matches expected ${args.expected_total:,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())